"""CSV/Excel bulk recipient import with validation and duplicate detection."""

import csv
import io
import re
import secrets
from typing import BinaryIO

from dropoutjeep_lib import db
from dropoutjeep_lib.models import Recipient
from dropoutjeep_lib.utils.suppression import is_suppressed

EMAIL_RE = re.compile(
    r"^[a-zA-Z0-9.!#$%&'*+/=?^_`{|}~-]+@[a-zA-Z0-9](?:[a-zA-Z0-9-]{0,61}[a-zA-Z0-9])?(?:\.[a-zA-Z0-9](?:[a-zA-Z0-9-]{0,61}[a-zA-Z0-9])?)*$"
)

DEFAULT_MAX_ROWS = int(__import__('os').environ.get('BULK_IMPORT_MAX_ROWS', '10000'))


def _effective_max_rows(max_rows: int | None) -> int:
    """Return explicit max_rows or the configured default."""
    if max_rows is not None:
        return max_rows
    from dropoutjeep_lib.utils.settings import get_setting_int
    return get_setting_int('bulk_import_max_rows', DEFAULT_MAX_ROWS)


def _normalize_row(row: dict) -> dict:
    """Return a normalized dict with lowercase keys and stripped string values."""
    out = {}
    for key, value in row.items():
        if key is None:
            continue
        k = str(key).strip().lower()
        if value is None:
            out[k] = ''
        else:
            out[k] = str(value).strip()
    return out


def _read_rows_from_file(file_obj: BinaryIO, filename: str) -> list[dict]:
    """Read rows from a CSV or XLSX file into normalized dicts."""
    lower_name = (filename or '').lower()
    if lower_name.endswith('.xlsx'):
        return _read_xlsx(file_obj)
    return _read_csv(file_obj)


def _read_csv(file_obj: BinaryIO) -> list[dict]:
    # CSV files are text; decode as utf-8-sig to handle BOM.
    text = io.TextIOWrapper(file_obj, encoding='utf-8-sig', newline='')
    reader = csv.DictReader(text)
    return [_normalize_row(row) for row in reader]


def _read_xlsx(file_obj: BinaryIO) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError('openpyxl is required for Excel import') from exc
    wb = load_workbook(file_obj, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip().lower() if h else '' for h in rows[0]]
    result = []
    for raw in rows[1:]:
        row = {headers[i]: (str(v).strip() if v is not None else '') for i, v in enumerate(raw)}
        result.append(row)
    return result


def _validate_email(email: str) -> str | None:
    """Return an error message if the email is invalid, else None."""
    if not email:
        return 'email is required'
    if len(email) > 320:
        return 'email exceeds 320 characters'
    if not EMAIL_RE.match(email):
        return 'invalid email format'
    local, _, domain = email.rpartition('@')
    if not local or not domain or '.' not in domain:
        return 'invalid email domain'
    return None


def import_recipients(
    file_obj: BinaryIO,
    filename: str,
    campaign_id: int | None = None,
    import_batch_id: str | None = None,
    max_rows: int | None = None,
    skip_duplicates: bool = True,
    global_duplicate_check: bool = False,
):
    """Import recipients from CSV/Excel.

    Returns a dict:
      - created: int
      - duplicates: int
      - errors: list[dict] with row, email, reason
      - import_batch_id: str
    """
    if import_batch_id is None:
        import_batch_id = secrets.token_urlsafe(12)

    effective_max = _effective_max_rows(max_rows)
    rows = _read_rows_from_file(file_obj, filename)
    if len(rows) > effective_max:
        return {
            'created': 0,
            'duplicates': 0,
            'errors': [{'row': 0, 'email': '', 'reason': f'file exceeds max {effective_max} rows'}],
            'import_batch_id': import_batch_id,
        }

    errors = []
    created = 0
    duplicates = 0
    seen_in_file = set()

    existing_emails = set()
    if skip_duplicates:
        query = Recipient.query
        if campaign_id is not None and not global_duplicate_check:
            query = query.filter_by(campaign_id=campaign_id)
        for r in query.with_entities(Recipient.email).all():
            existing_emails.add(r.email.lower())

    for idx, row in enumerate(rows, start=2):  # row 1 is header
        email = row.get('email', '')
        name = row.get('name', '')
        external_id = row.get('external_id', '')

        err = _validate_email(email)
        if err:
            errors.append({'row': idx, 'email': email, 'reason': err})
            continue

        if is_suppressed(email):
            errors.append({'row': idx, 'email': email, 'reason': 'suppressed'})
            continue

        email_lower = email.lower()
        if email_lower in seen_in_file:
            duplicates += 1
            continue
        seen_in_file.add(email_lower)

        if skip_duplicates and email_lower in existing_emails:
            duplicates += 1
            continue

        recipient = Recipient(
            email=email,
            name=name or None,
            external_id=external_id or None,
            campaign_id=campaign_id,
            import_batch_id=import_batch_id,
        )
        db.session.add(recipient)
        created += 1

    db.session.commit()
    return {
        'created': created,
        'duplicates': duplicates,
        'errors': errors,
        'import_batch_id': import_batch_id,
    }
